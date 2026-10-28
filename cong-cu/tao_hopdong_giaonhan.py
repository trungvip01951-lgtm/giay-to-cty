#!/usr/bin/env python3
"""Hợp đồng mua bán (thanh toán 100%) và Biên bản giao nhận cho một đơn hàng đã xuất hóa đơn.
Số liệu đọc từ JSON của đơn (mẫu: giay-to-khach-hang/TienDai/don_HD00002469.json); các điều khoản chuẩn lấy thẳng
từ mục 3.5 của hồ sơ (ho-so/HoSo_...md), không gõ lại. Số, ngày hợp đồng và ngày biên bản để trống cho anh điền.
Chạy: python3 cong-cu/tao_hopdong_giaonhan.py don.json [anh_dau.png]  -> ghi cạnh tệp JSON (bản không dấu; có ảnh thì thêm _CoDau)."""
import json
import pathlib
import re
import sys

sys.path.insert(0, str(pathlib.Path(__file__).resolve().parent))
import tao_mau as m  # noqa: E402
from tao_denghi_fata import doc_so, so  # noqa: E402
from docx.shared import Pt  # noqa: E402

TK, DD, BANK = m.TK, m.DD, m.BANK
HOSO = m.ROOT / "ho-so" / "HoSo_ThongTinCongTy_VaMauGiayTo_TrungKien.md"


def dieu_khoan_chuan():
    """Trả về (các dòng căn cứ, các dòng từ Điều 3 đến hết) của mẫu 3.5 trong hồ sơ."""
    t = HOSO.read_text(encoding="utf-8")
    khoi = t[t.index("### 3.5"):]
    khoi = khoi[khoi.index("```\n") + 4:]
    khoi = khoi[:khoi.index("```")]
    dong = [x.rstrip() for x in khoi.split("\n")]
    can_cu = [x[2:].strip() for x in dong if x.startswith("- Căn cứ")]
    i = next(k for k, x in enumerate(dong) if x.startswith("Điều 3:"))
    j = next(k for k, x in enumerate(dong) if x.startswith("| ĐẠI DIỆN BÊN MUA"))
    return can_cu, [x.strip() for x in dong[i:j] if x.strip()]


def thong_so(don):
    h = don["hang"]
    tien = sum(x["sl"] * x["don_gia"] for x in h)
    vat = tien * don["vat_phan_tram"] // 100
    tong = tien + vat
    k = don["kiem_tra_hoa_don"]
    assert (tien, vat, tong) == (k["cong_tien_hang"], k["tien_thue"], k["tong"]), "Số tiền không khớp hóa đơn"
    return tien, vat, tong


def ben_mua(d, kh, label, goi):
    m.mixed(d, [(f"{label}: ", True), (kh["ten"], True), (f" ({goi})", False)])
    m.para(d, f"Địa chỉ: {kh['dia_chi']}", after=0)
    m.para(d, f"Điện thoại kế toán: {kh['ke_toan']}", after=0)
    m.para(d, f"Mã số thuế: {kh['mst']}", after=0)
    m.para(d, f"Số tài khoản: {kh['stk']} tại {kh['ngan_hang']}", after=0)
    m.para(d, f"Đại diện: {kh['dai_dien']}          Chức vụ: {kh['chuc_vu']}", after=6)


def hop_dong(don, anh_dau, out):
    kh, hd = don["khach"], don["hoa_don"]
    tien, vat, tong = thong_so(don)
    can_cu, dieu = dieu_khoan_chuan()
    d = m.new_doc()
    m.ANH_DAU = anh_dau
    m.quoc_hieu(d)
    m.title(d, "HỢP ĐỒNG KINH TẾ", "(Số: " + m.dots(30) + ")")
    for c in can_cu:
        m.para(d, "- " + c, italic=True, after=1)
    m.para(d, f"Hôm nay, ngày ..... tháng ..... năm 2026, chúng tôi đại diện cho các bên ký hợp đồng, gồm có:", after=4)
    ben_mua(d, kh, "BÊN MUA", "gọi tắt là bên A")
    m.ben_tk(d, "BÊN BÁN", "gọi tắt là bên B")
    m.para(d, "Điều 1: Nội dung Hợp đồng", bold=True, after=2, keep=True)
    m.para(d, "Bên A đồng ý mua, bên B đồng ý bán hàng hóa với đơn giá, số lượng, thông số kỹ thuật như sau:", after=3)
    rows = [[str(i), x["ten"], x["dvt"], str(x["sl"]), so(x["don_gia"]), so(x["sl"] * x["don_gia"])]
            for i, x in enumerate(don["hang"], 1)]
    rows += [["", "Tổng:", "", "", "", so(tien)],
             ["", f"VAT ({don['vat_phan_tram']}%):", "", "", "", so(vat)],
             ["", "TỔNG CỘNG:", "", "", "", so(tong)]]
    m.table(d, ["STT", "Nội dung hàng hóa / Model", "ĐVT", "SL", "Đơn giá (VNĐ)", "Thành tiền (VNĐ)"], rows,
            [1.3, 6.3, 1.5, 1.4, 2.7, 2.8], size=12, center_cols=(0, 2, 3, 4, 5))
    m.para(d, f"(Bằng chữ: {doc_so(tong)}./.)", italic=True)
    m.para(d, f"Hàng hóa đã được bên B xuất Hóa đơn GTGT số {hd['so']}, ký hiệu {hd['ky_hieu']}, ngày {hd['ngay']}.")
    m.para(d, "Điều 2: Phương thức thanh toán", bold=True, after=2, keep=True)
    m.para(d, "Việc thanh toán của Hợp đồng được thực hiện bằng đồng tiền Việt Nam, bằng tiền mặt hoặc chuyển khoản. "
              "Bên A thanh toán cho bên B 100% giá trị hợp đồng trước khi nhận hàng hoặc ngay khi nhận hàng.")
    for x in dieu:
        if re.match(r"^Điều \d+:", x):
            m.para(d, x, bold=True, after=2, keep=True)
        elif x.startswith("Địa chỉ giao nhận chứng từ"):
            m.para(d, f"Địa chỉ giao nhận chứng từ: {kh['nhan_ho_so']}; hóa đơn điện tử gửi về {kh['email_hoa_don']}.")
        elif x.startswith("Địa điểm giao hàng"):
            m.para(d, "3.2. " + x, after=3)
        else:
            m.para(d, x, after=3)
    m.sign_block(d, "ĐẠI DIỆN BÊN MUA", "ĐẠI DIỆN BÊN BÁN", left_note="(Ký, ghi rõ họ tên, đóng dấu)",
                 right_note="(Ký, ghi rõ họ tên, đóng dấu)", right_name=DD["ho_ten"])
    p = out / f"HopDong_HD{hd['so']}_TienDai{'_CoDau' if anh_dau else ''}.docx"
    d.save(p)
    return p


def giao_nhan(don, anh_dau, out):
    kh, hd = don["khach"], don["hoa_don"]
    d = m.new_doc()
    m.ANH_DAU = None          # biên bản do nhân viên ký, không dùng dấu chữ ký Giám đốc
    t = d.add_table(rows=1, cols=2)
    t.autofit = False
    m.shade_borders(t, borders=False)
    a, b = t.rows[0].cells
    m.run(a.paragraphs[0], TK["ten"], bold=True, size=12)
    m.run(a.add_paragraph(), f"ĐT: {TK['dien_thoai']}", size=12)
    b.paragraphs[0].alignment = m.C
    m.run(b.paragraphs[0], "CỘNG HÒA XÃ HỘI CHỦ NGHĨA VIỆT NAM", bold=True, size=12)
    for s, bd in (("Độc lập - Tự do - Hạnh phúc", True), ("---------------***---------------", False),
                  ("Hà Nội, ngày ..... tháng ..... năm 2026", False)):
        q = b.add_paragraph()
        q.alignment = m.C
        m.run(q, s, bold=bd, italic=not bd, size=12)
    m.fix_layout(t, [7.6, 8.4])
    m.para(d, "", after=4)
    m.title(d, "BIÊN BẢN GIAO NHẬN HÀNG HÓA")
    m.para(d, f"Căn cứ Hợp đồng kinh tế số {m.dots(20)} ngày ..... tháng ..... năm 2026 và Hóa đơn GTGT số {hd['so']}, "
              f"ký hiệu {hd['ky_hieu']}, ngày {hd['ngay']}.", italic=True, after=4)
    m.mixed(d, [("Bên giao: ", True), (TK["ten"], True)])
    m.para(d, f"Địa chỉ: {TK['dia_chi']}", after=0)
    m.para(d, f"Người đại diện: {m.dots(30)}   Chức vụ: {m.dots(20)}", after=5)
    m.mixed(d, [("Bên nhận: ", True), (kh["ten"], True)])
    m.para(d, f"Địa chỉ: {kh['dia_chi']}", after=0)
    m.para(d, f"Người đại diện: {m.dots(30)}   Chức vụ: {m.dots(20)}", after=5)
    m.para(d, "Cùng tiến hành giao nhận hàng hóa sau:", after=3)
    rows = [[str(i), x["ten"], x["dvt"], str(x["sl"])] for i, x in enumerate(don["hang"], 1)]
    m.table(d, ["STT", "Tên hàng hóa / Model", "ĐVT", "Số lượng"], rows, [1.4, 9.2, 2.2, 3.0], size=12, center_cols=(0, 2, 3))
    m.para(d, "Chất lượng: Hàng mới 100%, nguyên đai kiện, đủ số lượng như trên.", after=1)
    m.para(d, "Bảo hành: 12 tháng kể từ ngày ký biên bản này, theo điều kiện của Hợp đồng.", after=1)
    m.para(d, f"Địa điểm giao nhận: {m.dots(60)}", after=3)
    m.para(d, "Hai bên đã kiểm tra, xác nhận bàn giao đúng số lượng, chủng loại như trên. Biên bản được lập thành 02 bản, mỗi bên giữ 01 bản.", after=8)
    m.sign_block(d, "ĐẠI DIỆN BÊN GIAO", "ĐẠI DIỆN BÊN NHẬN", left_note="(Ký, ghi rõ họ tên)", right_note="(Ký, ghi rõ họ tên)",
                 left_name="Hà Thu Thảo")
    p = out / f"BienBanGiaoNhan_HD{hd['so']}_TienDai.docx"
    d.save(p)
    return p


if __name__ == "__main__":
    f = pathlib.Path(sys.argv[1])
    don = json.load(open(f, encoding="utf-8"))
    for anh in (None, *sys.argv[2:3]):
        print(hop_dong(don, anh, f.parent))
    print(giao_nhan(don, None, f.parent))
